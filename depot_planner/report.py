# -*- coding: utf-8 -*-
"""結果の書き出し(JSON / CSV / docx)。

docx の体裁は route_calcurate で手直しした報告書に合わせている:
見出し行は濃紺(365F91)に白字、行は区分(group)ごとに淡色で塗り分け、合計行は灰色。
"""

import csv
import io
import json
import re
import time
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from . import config
from .runner import route_text

HEADER_FILL = "365F91"
TOTAL_FILL = "D9D9D9"
GROUP_FILLS = ["DCE6F1", "FDE9D9", "E5E0EC", "F2DCDB", "EAF1DD", "FFF2CC", "DAEEF3", "FBD4B4"]
PROFILE_LABEL = {"driving": "自動車", "driving-traffic": "自動車(交通状況考慮)",
                 "walking": "徒歩", "cycling": "自転車"}
OBJECTIVE_LABEL = {"makespan": "最も遅い車両の所要時間を最小化",
                   "total": "全車両の総所要時間を最小化"}


def run_dir(result, base=config.RUNS_DIR):
    stamp = time.strftime("%Y%m%d_%H%M%S")
    name = re.sub(r'[\\/:*?"<>|\s]+', "_", result["scenario_name"])[:40]
    d = Path(base) / f"{stamp}_{name}"
    d.mkdir(parents=True, exist_ok=True)
    return d


def to_json(result):
    return json.dumps(result, ensure_ascii=False, indent=2)


def to_csv(result):
    """目的地ごとの担当(Excel で開けるよう BOM 付き UTF-8)。"""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["目的地ID", "目的地", "緯度", "経度", "担当補給地点ID", "担当補給地点",
                "移動(分)", "所要(分)"])
    for a in result["assignments"]:
        w.writerow([a["target_id"], a["target"], a["lat"], a["lon"],
                    a["depot_id"], a["depot"], a["travel_min"], a["time_min"]])
    return "﻿" + buf.getvalue()


def to_routes_csv(result):
    """車両ごとの立ち寄り順(1 行 = 1 立ち寄り)。ルート最適化をしていなければ None。"""
    rt = result.get("routing")
    if not rt:
        return None
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["車両", "待機拠点", "順番", "種別", "地点ID", "地点", "個数", "完了(分)",
                "緯度", "経度"])
    for v in rt["vehicles"]:
        for n, st in enumerate(v["stops"], 1):
            w.writerow([v["no"], v["start"], n, "積込" if st["kind"] == "pickup" else "配送",
                        st["id"], st["name"], st.get("count", 1), st["done_min"],
                        st["lat"], st["lon"]])
    return "﻿" + buf.getvalue()


# ---------------------------------------------------------------------------
# docx
# ---------------------------------------------------------------------------
def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), fill)
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _table(doc, header, rows, fills=None, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = str(text)
        _shade(cell, HEADER_FILL)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # 見出し行を各ページで繰り返す
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = "" if text is None else str(text)
            if fills and fills[n]:
                _shade(cell, fills[n])
    if widths:
        for row in t.rows:
            for cell, wdt in zip(row.cells, widths):
                cell.width = Pt(wdt)
    return t


def to_docx(result, path):
    s, p = result["summary"], result["params"]
    groups = []
    for c in result["candidates"]:
        if c["group"] not in groups:
            groups.append(c["group"])
    fill_of = {g: GROUP_FILLS[i % len(GROUP_FILLS)] for i, g in enumerate(groups)}
    group_of = {c["id"]: c["group"] for c in result["candidates"]}

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Yu Gothic"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    st.font.size = Pt(10.5)

    title = doc.add_heading(f"補給地点選定シミュレーション（{result['scenario_name']}）", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("0. 前提条件", level=1)
    for line in [
        f"目的地 {s['num_targets']} 件／補給地点 候補 {s['num_candidates']} 件"
        f"（除外 {s['num_excluded']} 件）から {s['num_depots']} か所を選定。",
        f"車両 1 台に積める機材 {s['vehicle_capacity']} 個 → 車両 {s['num_vehicles']} 台。",
        f"移動時間: {result['provider']}（{PROFILE_LABEL.get(p['profile'], p['profile'])}、道路ネットワーク上の所要時間）。",
        f"作業時間: 目的地での作業 {p['service_min']} 分／補給地点での積込 {p['reload_min']} 分。",
        f"目的関数: {OBJECTIVE_LABEL[p['objective']]}。",
        f"探索: {'全組合せの総当たり（厳密解）' if s['method'] == 'exact' else '貪欲法＋入替え探索（近似解）'}"
        f"。組合せ数 {s['combinations']:,}。",
    ]:
        doc.add_paragraph(line, style="List Bullet")

    doc.add_heading("1. 補給地点 候補リスト（全件）", level=1)
    _table(doc, ["採否", "区分", "名称", "緯度", "経度", "所在地"],
           [[c["result"], c["group"] or "-", c["name"], f"{c['lat']:.6f}", f"{c['lon']:.6f}",
             c["address"] or "-"] for c in result["candidates"]],
           fills=[fill_of[c["group"]] for c in result["candidates"]],
           widths=[40, 62, 118, 58, 62, 140])

    doc.add_heading("2. シミュレーション結果：採用した補給地点と車両配分", level=1)
    para = doc.add_paragraph()
    run = para.add_run(f"最長所要時間{s['makespan_min']}分")
    run.bold, run.underline, run.font.size = True, True, Pt(14)
    rows = [[d["name"], d["group"] or "-", f"{d['stock']} 個", f"{d['vehicles']} 台",
             d["max_min"], ", ".join(d["targets"])] for d in result["depots"]]
    rows.append(["合計", "", f"{s['num_targets']} 個", f"{s['num_vehicles']} 台",
                 s["makespan_min"], ""])
    _table(doc, ["拠点", "区分", "機材", "配置台数", "最長所要(分)", "担当地点"], rows,
           fills=[fill_of[group_of[d["id"]]] for d in result["depots"]] + [TOTAL_FILL],
           widths=[116, 56, 40, 50, 56, 162])
    doc.add_paragraph(f"総所要 {s['total_min']} 分／目的地 1 件あたり平均 {s['avg_min']} 分"
                      "（各目的地へ補給地点から直行する前提の値）。"
                      f"拠点数を増やしても最長所要は {s['lower_bound_min']} 分を下回らない"
                      f"（{s['bottleneck_target']} が最寄りの補給地点からも遠いため）。")

    doc.add_heading("3. 採用拠点の所在地・メモ", level=1)
    _table(doc, ["拠点", "所在地", "メモ"],
           [[d["name"], d["address"] or "-", d["note"] or "-"] for d in result["depots"]],
           fills=[fill_of[group_of[d["id"]]] for d in result["depots"]],
           widths=[130, 170, 180])

    doc.add_heading("4. 目的地ごとの担当拠点（所要時間の遅い順）", level=1)
    _table(doc, ["順位", "目的地", "担当拠点", "移動(分)", "所要(分)"],
           [[i, a["target"], a["depot"], a["travel_min"], a["time_min"]]
            for i, a in enumerate(result["assignments"], 1)],
           widths=[34, 120, 180, 60, 60])

    if len(result["sensitivity"]) > 1:
        doc.add_heading("5. 採用地点数を変えた場合（感度分析）", level=1)
        _table(doc, ["拠点数", "最長所要(分)", "総所要(分)", "採用拠点（機材数）"],
               [[r["k"], r["makespan_min"], r["total_min"],
                 " / ".join(f"{d['name']}({d['stock']})" for d in r["depots"])]
                for r in result["sensitivity"]],
               fills=[TOTAL_FILL if r["k"] == s["num_depots"] else None
                      for r in result["sensitivity"]],
               widths=[40, 60, 60, 320])

    rt = result.get("routing")
    if rt:
        doc.add_heading("6. 配分ルート最適化：車両ごとの順路", level=1)
        para = doc.add_paragraph()
        run = para.add_run(f"最長所要時間{rt['makespan_min']}分")
        run.bold, run.underline, run.font.size = True, True, Pt(14)
        doc.add_paragraph(
            f"車両 {rt['num_vehicles']} 台（1 台に機材 {rt['vehicle_capacity']} 個まで、"
            f"1 台あたり最大 {rt['max_deliveries']} 件）。OR-Tools で最も遅い車両の所要時間を"
            f"最小化（探索 {rt['elapsed_sec']} 秒）。補給地点への帰路は含まない。")
        _table(doc, ["車両", "待機拠点", "配送", "所要(分)", "順路"],
               [[v["no"], v["start"], f"{v['deliveries']} 件", v["time_min"],
                 route_text(v).split(" → ", 1)[1]] for v in rt["vehicles"]],
               fills=[fill_of[group_of[v["start_id"]]] for v in rt["vehicles"]],
               widths=[34, 110, 40, 50, 246])

    if result["warnings"]:
        doc.add_heading("注意事項", level=1)
        for w in result["warnings"]:
            doc.add_paragraph(w, style="List Bullet")

    doc.save(path)
    return path


def save_all(result, directory=None, docx=True):
    """runs/<日時_名前>/ に result.json / assignments.csv / routes.csv / report.docx を保存する。"""
    d = Path(directory) if directory else run_dir(result)
    d.mkdir(parents=True, exist_ok=True)
    (d / "result.json").write_text(to_json(result), encoding="utf-8")
    (d / "assignments.csv").write_text(to_csv(result), encoding="utf-8")
    routes = to_routes_csv(result)
    if routes:
        (d / "routes.csv").write_text(routes, encoding="utf-8")
    if docx:
        to_docx(result, d / "report.docx")
    return d
