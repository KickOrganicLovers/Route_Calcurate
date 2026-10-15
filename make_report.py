# -*- coding: utf-8 -*-
"""発表用レポート(docx)生成。

2 つの表だけを出す。
    表1: 採用した補給地点リスト(配置台数・最長所要時間つき)
    表2: 補給地点の候補リスト全件(採用・不採用の別つき)

実行:
    uv run make_report.py
"""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from travel_time import build_travel_time
from main import build_inputs, run_selection
from waypoints import candidate_supply_points


def add_table(doc, headers, rows, style="Light Grid Accent 1"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = style
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    return t


def main():
    # ---- 計算 ----
    targets, candidates = build_inputs()
    travel_time_min = build_travel_time(targets, candidates, profile="driving")
    sel = run_selection(targets, candidates, travel_time_min)
    if sel is None:
        raise SystemExit("補給地点を選定できませんでした(条件を確認してください)")
    depots = sel["depots"]

    # 拠点ごとの最長所要時間(その拠点を使う車両のうち最も遅いもの)
    worst = {}
    for a in sel["assignments"]:
        worst[a["depot"]] = max(worst.get(a["depot"], 0), a["time_min"])

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Yu Gothic"
    style.font.size = Pt(10.5)

    title = doc.add_heading("配分ルート最適化 シミュレーション結果", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph(
        f"測定局 {len(targets)} 地点 / 補給地点 {len(depots)} か所"
        f"（候補 {len(candidate_supply_points)} 件から選定） / "
        f"車両 {sel['num_vehicles']} 台（1局1台）　"
        f"メイクスパン {sel['makespan_min']} 分")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ===== 表1: 補給地点リスト =====
    doc.add_heading("1. 補給地点リスト", level=1)
    add_table(doc,
              ["拠点", "配置台数", "最長所要(分)", "所在地", "集合駅",
               "レンタカー営業所"],
              [[f"[{d['code']}] {d['name']}", f"{d['stock']} 台",
                worst.get(d["name"], "-"), d.get("address", ""),
                d.get("station", "要確認"), d.get("rental", "要確認")]
               for d in depots])

    # ===== 表2: 候補地全件（採用・不採用の別つき） =====
    doc.add_heading("2. 補給地点 候補リスト（全件）", level=1)
    adopted = {d["code"] for d in depots}
    add_table(doc,
              ["code", "採否", "区分", "方面", "名称", "緯度", "経度", "所在地"],
              [[c["code"], "採用" if c["code"] in adopted else "-",
                c["kind"], c["area"], c["name"],
                f"{c['lat']:.6f}", f"{c['lon']:.6f}", c.get("address", "-")]
               for c in candidate_supply_points])

    out = "docs/配分ルート最適化_レポート_車両23台.docx"
    doc.save(out)
    print(f"出力: {out}")


if __name__ == "__main__":
    main()
